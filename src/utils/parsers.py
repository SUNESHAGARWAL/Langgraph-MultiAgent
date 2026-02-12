"""
Document Parsers for RAG Ingestion

Supports parsing multiple document formats:
- Plain text (.txt)
- Markdown (.md)
- PDF (.pdf)
- Word documents (.docx)
- CSV (.csv)
- Excel (.xlsx)

Features:
- Metadata extraction (filename, path, size, modified date)
- Text chunking with overlap
- Sentence boundary preservation
"""

import os
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)


@dataclass
class Document:
    """Document container for RAG"""
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    embedding: Optional[List[float]] = None


class DocumentParser:
    """
    Multi-format document parser for RAG ingestion

    Extracts text content and metadata from various document formats.
    """

    SUPPORTED_EXTENSIONS = {'.txt', '.md', '.pdf', '.docx', '.csv', '.xlsx'}

    def __init__(self):
        self._check_dependencies()

    def _check_dependencies(self):
        """Check if optional dependencies are available"""
        self.has_pdf = False
        self.has_docx = False
        self.has_pandas = False

        try:
            import pypdf
            self.has_pdf = True
        except ImportError:
            logger.warning("pypdf not installed - PDF parsing disabled")

        try:
            import docx
            self.has_docx = True
        except ImportError:
            logger.warning("python-docx not installed - DOCX parsing disabled")

        try:
            import pandas
            self.has_pandas = True
        except ImportError:
            logger.warning("pandas not installed - CSV/Excel parsing disabled")

    def parse(self, file_path: str) -> Document:
        """
        Parse document from file

        Args:
            file_path: Path to document file

        Returns:
            Document with content and metadata

        Raises:
            FileNotFoundError: If file doesn't exist
            ValueError: If file format not supported
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        extension = path.suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file format: {extension}")

        # Extract metadata
        metadata = self._extract_metadata(path)

        # Parse content based on extension
        try:
            if extension == '.txt':
                content = self._parse_txt(path)
            elif extension == '.md':
                content = self._parse_md(path)
            elif extension == '.pdf':
                content = self._parse_pdf(path)
            elif extension == '.docx':
                content = self._parse_docx(path)
            elif extension == '.csv':
                content = self._parse_csv(path)
            elif extension == '.xlsx':
                content = self._parse_excel(path)
            else:
                raise ValueError(f"Parser not implemented for {extension}")

            logger.info(f"✓ Parsed {path.name} ({len(content)} characters)")

            return Document(content=content, metadata=metadata)

        except Exception as e:
            logger.error(f"Failed to parse {path.name}: {e}", exc_info=True)
            raise

    def _extract_metadata(self, path: Path) -> Dict[str, Any]:
        """Extract file metadata"""
        stats = path.stat()
        return {
            "filename": path.name,
            "file_path": str(path.absolute()),
            "extension": path.suffix.lower(),
            "size_bytes": stats.st_size,
            "modified": datetime.fromtimestamp(stats.st_mtime).isoformat(),
            "created": datetime.fromtimestamp(stats.st_ctime).isoformat(),
        }

    def _parse_txt(self, path: Path) -> str:
        """Parse plain text file"""
        with open(path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    def _parse_md(self, path: Path) -> str:
        """Parse markdown file (treat as plain text)"""
        return self._parse_txt(path)

    def _parse_pdf(self, path: Path) -> str:
        """Parse PDF file"""
        if not self.has_pdf:
            raise ImportError("pypdf not installed. Install with: pip install pypdf")

        import pypdf

        text_parts = []

        try:
            with open(path, 'rb') as f:
                pdf_reader = pypdf.PdfReader(f)

                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        text = page.extract_text()
                        if text.strip():
                            text_parts.append(f"[Page {page_num + 1}]\n{text}")
                    except Exception as e:
                        logger.warning(f"Failed to extract page {page_num + 1}: {e}")

        except Exception as e:
            logger.error(f"Failed to read PDF: {e}")
            raise

        return "\n\n".join(text_parts)

    def _parse_docx(self, path: Path) -> str:
        """Parse Word document"""
        if not self.has_docx:
            raise ImportError("python-docx not installed. Install with: pip install python-docx")

        import docx

        doc = docx.Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract table content
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    table_texts.append(row_text)

        all_text = "\n\n".join(paragraphs)
        if table_texts:
            all_text += "\n\n[Tables]\n" + "\n".join(table_texts)

        return all_text

    def _parse_csv(self, path: Path) -> str:
        """Parse CSV file"""
        if not self.has_pandas:
            raise ImportError("pandas not installed. Install with: pip install pandas")

        import pandas as pd

        df = pd.read_csv(path)

        # Convert to text representation
        content = f"CSV File: {path.name}\n"
        content += f"Columns: {', '.join(df.columns)}\n"
        content += f"Rows: {len(df)}\n\n"

        # Include first few rows as sample
        content += "Sample Data:\n"
        content += df.head(10).to_string(index=False)

        # Include summary statistics for numeric columns
        numeric_cols = df.select_dtypes(include=['number']).columns
        if len(numeric_cols) > 0:
            content += "\n\nSummary Statistics:\n"
            content += df[numeric_cols].describe().to_string()

        return content

    def _parse_excel(self, path: Path) -> str:
        """Parse Excel file"""
        if not self.has_pandas:
            raise ImportError("pandas not installed. Install with: pip install pandas openpyxl")

        import pandas as pd

        # Read all sheets
        excel_file = pd.ExcelFile(path)
        content_parts = [f"Excel File: {path.name}"]
        content_parts.append(f"Sheets: {', '.join(excel_file.sheet_names)}\n")

        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(path, sheet_name=sheet_name)

            content_parts.append(f"\n[Sheet: {sheet_name}]")
            content_parts.append(f"Columns: {', '.join(df.columns)}")
            content_parts.append(f"Rows: {len(df)}\n")

            # Sample data
            content_parts.append("Sample Data:")
            content_parts.append(df.head(5).to_string(index=False))

        return "\n".join(content_parts)

    @staticmethod
    def chunk_text(
        text: str,
        chunk_size: int = 1000,
        chunk_overlap: int = 200,
        separator: str = "\n\n"
    ) -> List[str]:
        """
        Split text into chunks with overlap

        Args:
            text: Input text to chunk
            chunk_size: Maximum characters per chunk
            chunk_overlap: Overlap between chunks
            separator: Separator to split on (tries to preserve paragraphs)

        Returns:
            List of text chunks
        """
        if not text:
            return []

        if len(text) <= chunk_size:
            return [text]

        chunks = []

        # Try to split on separator first (preserve paragraphs)
        paragraphs = text.split(separator)

        current_chunk = ""
        for para in paragraphs:
            # If adding this paragraph exceeds chunk size
            if len(current_chunk) + len(para) + len(separator) > chunk_size:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                    # Start new chunk with overlap
                    current_chunk = current_chunk[-chunk_overlap:] + separator + para
                else:
                    # Paragraph itself is too large - split it
                    if len(para) > chunk_size:
                        # Split on sentences
                        sentences = para.split('. ')
                        for sentence in sentences:
                            if len(current_chunk) + len(sentence) > chunk_size:
                                if current_chunk:
                                    chunks.append(current_chunk.strip())
                                    current_chunk = current_chunk[-chunk_overlap:] + '. ' + sentence
                                else:
                                    # Even sentence is too large - hard split
                                    for i in range(0, len(sentence), chunk_size - chunk_overlap):
                                        chunk = sentence[i:i + chunk_size]
                                        chunks.append(chunk)
                            else:
                                current_chunk += '. ' + sentence
                    else:
                        current_chunk = para
            else:
                if current_chunk:
                    current_chunk += separator + para
                else:
                    current_chunk = para

        # Add final chunk
        if current_chunk:
            chunks.append(current_chunk.strip())

        logger.info(f"Split text into {len(chunks)} chunks (size: {chunk_size}, overlap: {chunk_overlap})")

        return chunks


def get_document_parser() -> DocumentParser:
    """Get document parser instance"""
    return DocumentParser()


def load_documents_from_directory(
    directory_path: str,
    recursive: bool = True,
    extensions: Optional[List[str]] = None
) -> List[Document]:
    """
    Load all documents from directory

    Args:
        directory_path: Path to directory
        recursive: Search subdirectories
        extensions: File extensions to include (default: all supported)

    Returns:
        List of parsed documents
    """
    path = Path(directory_path)

    if not path.exists():
        raise FileNotFoundError(f"Directory not found: {directory_path}")

    parser = DocumentParser()
    documents = []

    # Get all files
    if recursive:
        files = path.rglob('*')
    else:
        files = path.glob('*')

    # Filter by extension
    if extensions is None:
        extensions = list(parser.SUPPORTED_EXTENSIONS)
    else:
        extensions = [ext if ext.startswith('.') else f'.{ext}' for ext in extensions]

    # Parse each file
    for file_path in files:
        if file_path.is_file() and file_path.suffix.lower() in extensions:
            try:
                doc = parser.parse(str(file_path))
                documents.append(doc)
            except Exception as e:
                logger.error(f"Failed to parse {file_path.name}: {e}")

    logger.info(f"✓ Loaded {len(documents)} documents from {directory_path}")

    return documents
